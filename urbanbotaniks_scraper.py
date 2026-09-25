#!/usr/bin/env python3
"""UrbanBotaniks category 62 -> the supplied Temu pots template."""
from __future__ import annotations

import argparse
import concurrent.futures
import json
import logging
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

from lxml import html
from openpyxl import load_workbook

ROOT = "https://www.urbanbotaniks.com"
CATEGORY = ROOT + "/category/62/saksii-i-podlozhki.html"
LOG = logging.getLogger("urbanbotaniks")
EU = {"Austria", "Belgium", "Bulgaria", "Croatia", "Cyprus", "Czech Republic", "Denmark", "Estonia", "Finland", "France", "Germany", "Greece", "Hungary", "Ireland", "Italy", "Latvia", "Lithuania", "Luxembourg", "Malta", "Netherlands", "Poland", "Portugal", "Romania", "Slovakia", "Slovenia", "Spain", "Sweden"}
COUNTRIES = {
    "италия": "Italy", "italy": "Italy", "турция": "Türkiye", "turkey": "Türkiye",
    "румъния": "Romania", "romania": "Romania", "испания": "Spain", "spain": "Spain",
    "българия": "Bulgaria", "bulgaria": "Bulgaria", "китай": "China", "china": "China",
    "холандия": "Netherlands", "нидерландия": "Netherlands", "netherlands": "Netherlands",
    "германия": "Germany", "germany": "Germany", "франция": "France", "france": "France",
}
MANUFACTURERS = {
    "nuova pasquini": ("Nuova Pasquini & Bini S.p.A.", "Italy"),
    "p&b": ("Nuova Pasquini & Bini S.p.A.", "Italy"),
    "vega black": ("Nuova Pasquini & Bini S.p.A.", "Italy"),
    "gronest": ("Gronest Europe Srl", "Romania"),
    "vanguard": ("Vanguard hydroponics", "Spain"),
    "senkap": ("Senkap Ambalaj Plast.Tarm.Elk.Hırd.Ltd.Şti.", "Türkiye"),
}
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; UrbanBotaniksCatalog/1.0)", "Accept-Language": "bg-BG,bg;q=0.9,en;q=0.7"}


def clean(text):
    return re.sub(r"\s+", " ", text or "").strip()


def first(node, query, default=""):
    found = node.xpath(query)
    if not found:
        return default
    obj = found[0]
    return clean(obj if isinstance(obj, str) else obj.text_content())


def fetch(url):
    for attempt in range(4):
        try:
            req = urllib.request.Request(url, headers=HEADERS)
            with urllib.request.urlopen(req, timeout=45) as response:
                if urllib.parse.urlparse(response.url).netloc != "www.urbanbotaniks.com":
                    raise ValueError("Unexpected redirect: " + response.url)
                data = response.read()
                if len(data) < 3000:
                    raise ValueError("Unexpectedly short response")
                return html.fromstring(data, base_url=url)
        except (urllib.error.URLError, TimeoutError, ValueError) as exc:
            if attempt == 3:
                raise RuntimeError(f"Could not fetch {url}: {exc}") from exc
            time.sleep(2 ** attempt)


def category_links():
    """Follow pagination only inside the exact category; no parent/sibling categories."""
    pending = [CATEGORY]
    seen_pages, found = set(), {}
    while pending:
        url = pending.pop(0)
        if url in seen_pages:
            continue
        seen_pages.add(url)
        page = fetch(url)
        links = page.xpath('//a[contains(@class,"c-product-grid__product-title-link")]/@href')
        if not links:
            raise RuntimeError("No product cards on category page " + url)
        for link in links:
            target = urllib.parse.urljoin(ROOT, link)
            if re.match(r"https://www\.urbanbotaniks\.com/product/\d+/", target):
                found[target] = None
        for link in page.xpath('//a[contains(@class,"c-pager__page-number")]/@href'):
            target = urllib.parse.urljoin(ROOT, link)
            if urllib.parse.urlparse(target).path == urllib.parse.urlparse(CATEGORY).path and target not in seen_pages and target not in pending:
                pending.append(target)
        LOG.info("Category page %s: %d cards, %d distinct products", url, len(links), len(found))
    return list(found)


def parse_money(text):
    m = re.search(r"(?:€\s*)?(\d+(?:[.,]\d{1,2})?)", text or "")
    return float(m.group(1).replace(",", ".")) if m else None


def identify_country(text):
    low = text.casefold()
    # Use explicit production wording, never infer physical origin from a brand alone.
    for pattern in [r"(?:произведен[ао]?|производство|made in|произход)\s*(?:в|от|:)?\s*([А-Яа-яA-Za-z]+)", r"(?:страна на произход)\s*:\s*([А-Яа-яA-Za-z]+)"]:
        for match in re.finditer(pattern, low):
            country = COUNTRIES.get(match.group(1))
            if country:
                return country
    return ""


def parse_product(url):
    page = fetch(url)
    title = first(page, '//h1[contains(@class,"c-product-page__product-name")]')
    if not title:
        raise ValueError("No title at " + url)
    description_node = page.xpath('//*[contains(@class,"c-product-page__product-description")]')
    description = clean(description_node[0].text_content()) if description_node else ""
    brand = first(page, '//a[contains(@class,"c-product-page__product-brand-link")]')
    code = first(page, '//*[@id="ProductCode"]')
    ident = re.search(r"/product/(\d+)/", url).group(1)
    sku = re.sub(r"[^A-Za-z0-9-]", "-", code) if code else ident
    weight_s = first(page, '//*[@id="ProductWeight"]')
    weight_g = round(float(weight_s.replace(",", ".")) * 1000) if re.fullmatch(r"\d+(?:[.,]\d+)?", weight_s) else None

    price = first(page, '//*[contains(@class,"c-product-page__product-price-section")]//*[@itemprop="price"]/@content')
    if not price:
        price = first(page, '//*[contains(@class,"c-product-page__product-price-section")]//*[@itemprop="price"]')
    if not price:
        price = first(page, '//*[contains(@class,"c-product-page__product-price-section")]//*[contains(@class,"u-price__base__value")]/text()')
    sale = parse_money(price)
    if sale is None:
        raise ValueError("No EUR price at " + url)
    former = first(page, '//*[contains(@class,"c-product-page__product-price-section")]//*[contains(@class,"old-price") or contains(@class,"list-price__value")]/text()')
    list_price = parse_money(former)
    if list_price is not None and list_price <= sale:
        list_price = None
    available = page.xpath('//*[@itemprop="availability"]/@content')
    in_stock = not available or all("OutOfStock" not in x and "SoldOut" not in x for x in available)

    images = []
    for link in page.xpath('//a[contains(@class,"c-product-page__product-image-with-zoom") or contains(@class,"c-product-page__thumb-link")]/@href'):
        image = urllib.parse.urljoin(ROOT, link)
        if image.startswith(ROOT + "/userfiles/") and image not in images:
            images.append(image)
    if not images:
        raise ValueError("No product images at " + url)

    text = title + " " + description
    manufacturer = ""
    brand_key = brand.casefold()
    for token, (chosen, _) in MANUFACTURERS.items():
        if token in brand_key:
            manufacturer = chosen
            break
    # Recognizable model marks are useful if the site's brand metadata is absent.
    if not manufacturer:
        for token, (chosen, _) in MANUFACTURERS.items():
            if token in text.casefold():
                manufacturer = chosen
                break
    country = identify_country(description + " " + title)
    if not country and manufacturer:
        # Manufacturer headquarters is a clue for review, not evidence of production.
        LOG.warning("Origin unspecified for %s; manufacturer is %s", url, manufacturer)
    material = "Plastic" if re.search(r"пластмас|полиетилен|полипропилен|pvc", text, re.I) else ("Fabric" if re.search(r"текстил", text, re.I) else "")
    # Fabric is not a permitted Pots material; preserve evidence in description only.
    if material == "Fabric":
        material = ""
    shape = "Round" if re.search(r"кръгл|⌀", text, re.I) else ("Square" if re.search(r"квадратн", text, re.I) else "")
    # A saucer title may say "for a 5 L pot"; that is not the saucer's capacity.
    volume = re.search(r"вместимост\s*:\s*(\d+(?:[,.]\d+)?)\s*(?:л\.|литра|литри|l\b)", description, re.I)
    capacity = float(volume.group(1).replace(",", ".")) if volume else None
    dims = re.search(r"(\d+(?:[.,]\d+)?)\s*[xх×]\s*(\d+(?:[.,]\d+)?)(?:\s*[xх×]\s*(\d+(?:[.,]\d+)?))?\s*см", text, re.I)
    measures = [float(x.replace(",", ".")) for x in dims.groups() if x] if dims else []
    return dict(url=url, id=ident, sku="UB-" + sku, title=title, description=description, brand=brand,
                manufacturer=manufacturer, origin=country, weight_g=weight_g, price=sale, list_price=list_price,
                images=images[:12], in_stock=in_stock, material=material, shape=shape, capacity=capacity, dimensions=measures)


def choices(wb, category, suffix):
    key = f"t_8_{category}_{suffix}"
    row = next(r for r in wb["Dropdown Lists"] if r[0].value == key)
    return {str(c.value) for c in row[2:] if c.value not in (None, "None")}


def write_workbook(products, template, output, qty):
    wb = load_workbook(template)
    ws = wb["Template"]
    if ws["E4"].value != "t_1_Category" or ws["IS4"].value != "t_8_Governance Property:3":
        raise ValueError("Unexpected Temu template columns")
    if len(products) > 4996:
        raise ValueError("More products than template rows")
    for row in ws.iter_rows(min_row=5, max_row=min(5000, ws.max_row)):
        for cell in row:
            if cell.column != 6 and cell.value is not None:  # preserve F category lookup formula
                cell.value = None
    valid_manufacturer = choices(wb, 25104, "Manufacturer")
    valid_rp = choices(wb, 25104, "EU Responsible person")
    valid_origin = choices(wb, 25104, "Country/Region of Origin")
    issues = []
    for i, p in enumerate(products, 5):
        values = {"E": 25104, "G": "Normal product", "L": p["title"], "M": p["sku"], "N": p["sku"],
                  "R": p["brand"] or None, "T": p["description"] or p["title"],
                  "DS": p["shape"] or None, "DU": p["material"] or None,
                  "EU": "Model", "FG": p["sku"],
                  "FI": p["images"][0], "FT": qty if p["in_stock"] else 0,
                  "FU": p["price"], "FV": p["url"], "FW": p["list_price"],
                  "FX": "N/A" if p["list_price"] is None else None,
                  "FY": p["weight_g"], "GN": "UB Склад София", "GO": "1 Day",
                  "GP": "I will ship this item myself", "GR": p["origin"] or None,
                  "IR": p["sku"], "IS": p["manufacturer"] or None}
        if p["origin"] and p["origin"] not in EU:
            values["IT"] = "Shteryo Shterev"
        if p["capacity"] is not None:
            values.update({"EN": p["capacity"], "EO": "L"})
        if p["dimensions"]:
            # Product dimensions are not reliable package dimensions; leave FZ:GB for review.
            pass
        for offset, image in enumerate(p["images"][1:10], 1):
            values["FI FJ FK FL FM FN FO FP FQ FR".split()[offset]] = image
        for offset, image in enumerate(p["images"][:10]):
            values["AA AB AC AD AE AF AG AH AI AJ".split()[offset]] = image
        for col, val in values.items():
            ws[f"{col}{i}"] = val
        problem = []
        if p["manufacturer"] not in valid_manufacturer:
            problem.append("manufacturer unverified or absent")
        if p["origin"] not in valid_origin:
            problem.append("country of manufacture unverified or absent")
        if p["origin"] and p["origin"] not in EU and values.get("IT") not in valid_rp:
            problem.append("EU responsible person not in dropdown")
        if p["weight_g"] is None:
            problem.append("weight not stated")
        if not p["material"]:
            problem.append("material not stated/mapped")
        if not p["in_stock"]:
            problem.append("out of stock")
        if problem:
            issues.append(dict(row=i, url=p["url"], title=p["title"], issues=problem, brand=p["brand"], manufacturer=p["manufacturer"], origin=p["origin"]))
    wb.save(output)
    return issues


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", default="TEMU_POTS_STORAGE-TRUNKS.xlsx")
    parser.add_argument("--output", default="urbanbotaniks-temu.xlsx")
    parser.add_argument("--quantity", type=int, default=100, help="Assumed quantity for products shown as in stock; confirm with seller")
    parser.add_argument("--workers", type=int, default=6)
    parser.add_argument("--limit", type=int, default=0, help="Small test run")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    links = category_links()
    if args.limit:
        links = links[:args.limit]
    products, failed = [], []
    with concurrent.futures.ThreadPoolExecutor(max_workers=max(1, min(args.workers, 8))) as pool:
        tasks = {pool.submit(parse_product, u): u for u in links}
        for task in concurrent.futures.as_completed(tasks):
            try:
                products.append(task.result())
            except Exception as exc:
                LOG.error("%s", exc)
                failed.append({"url": tasks[task], "error": str(exc)})
            LOG.info("Processed %d/%d", len(products) + len(failed), len(links))
    products.sort(key=lambda p: int(p["id"]))
    issues = write_workbook(products, args.template, args.output, args.quantity)
    report = {"category": CATEGORY, "discovered": len(links), "written": len(products), "failed": failed,
              "review": issues, "quantity_assumption": args.quantity,
              "global_review": "Not ready for Temu import: package dimensions FZ:GB, packaging fields GD:GF, and indoor/outdoor usage ED require verified seller data. Quantity 100 is an assumption; weight is the product weight supplied by the retailer, not verified packaged weight.",
              "required_fields_pending_seller_input": ["ED Indoor Outdoor Usage", "FZ Package Length", "GA Package Width", "GB Package Height", "GD Individually packed", "GE Total packaging quantity", "GF Packaging unit"]}
    report_file = Path(args.output).with_suffix(".review.json")
    report_file.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    LOG.info("Saved %s and %s; %d review rows; %d failures", args.output, report_file, len(issues), len(failed))
    if failed or not products:
        raise SystemExit("Incomplete scrape; see review JSON")


if __name__ == "__main__":
    main()
